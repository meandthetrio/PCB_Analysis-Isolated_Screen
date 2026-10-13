#!/usr/bin/env python3
"""Build rev4/order/STOCK_CHECK.csv and rev4/order/Rev4_JLCPCB_Order.xlsx (BOM_SMT, Hand_Solder, Stock_Check)
from the CSVs made by make_order_bom.py and the JLCPCB parts API dump kicad_out/rev4/jlc_parts.json."""
import csv, json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
jlc = json.load(open('kicad_out/rev4/jlc_parts.json'))
bom = list(csv.DictReader(open('rev4/order/BOM_JLCPCB_SMT.csv')))
hand = list(csv.DictReader(open('rev4/order/HAND_SOLDER.csv')))
ALT = {  # LCSC -> (alternative, note); only where it matters or saves an Extended fee (verified via the JLCPCB parts API 2026-10-13)
 'C130212': ('none in JLC stock', 'CRITICAL: 5 in stock, min buy 7. TPA6111A2DR (C497472, 887 in stock) is SOIC-8, not the HVSSOP-8-EP footprint; TPA6111A2DGN(R) HVSSOP parts show 0 stock. Options: buy TPA6110A2DGNR from Mouser/Digi-Key and consign or hand-solder U3; or JLC global-sourcing pre-order (lead time); or change U3 to SOIC-8 on the next spin'),
 'C385198': ('C191023', 'optional: 1N5819WS SOD-323 40V 1A, Basic, 4.4M stock — drop-in for D1 if NSR1020MW2T3G runs short (501 in stock)'),
 'C1590':   ('C14663', 'optional Basic swap: CC0603KRX7R9BB104 100nF 50V X7R (saves one Extended fee)'),
 'C1226':   ('C22962', 'optional Basic swap: 0603WAF2200T5E 220R 1%'),
 'C25600':  ('C25819', 'optional Basic swap: 0603WAF4702T5E 47k 1%'),
 'C23344':  ('C31850', 'optional Basic swap: 0603WAF2202T5E 22k 1%'),
 'C1607':   ('C23630', 'optional Basic swap: CL10A225KO8NNNC 2.2uF 16V X5R'),
 'C1705':   ('C19666', 'optional Basic swap: CL10A475KO8NNNC 4.7uF 16V X5R'),
 'C14675':  ('C25803', 'optional Basic swap: 0603WAF1003T5E 100k 1%'),
 'C22369937':('C25804', 'optional Basic swap: 0603WAF1002T5E 10k 1%'),
 'C14676':  ('C21190', 'optional Basic swap: 0603WAF1001T5E 1k 1%'),
}
LEAD = {  # lead-time / sourcing notes, 2026-10-13 (JLCPCB help pages; TI product page; Octopart snapshot via search for distributor stock)
 'C130212': 'SHORT at JLCPCB (5 pcs, min buy 7; LCSC 5 pcs, pre-sale flagged; TI store: out of stock; lifecycle ACTIVE). Fastest: Digi-Key (~528 pcs) or Mouser (~2,165 pcs, aggregator snapshot, verify live) ship in-stock parts in 1-3 days + transit, then hand-solder or consign. JLCPCB global sourcing: 1-3 business days review + 9-20 business days, assembly waits for arrival. JLCPCB pre-order: no lead time published.',
 'C385198': 'In stock at JLCPCB (501, 1 per board): no part lead time for runs up to ~450 boards after attrition. Beyond stock: global sourcing 9-20 business days, or Basic drop-in C191023 (4.4M pcs).',
 'C114916': 'In stock at JLCPCB (1,407): no part lead time. Beyond stock: global sourcing 9-20 business days.',
 'C53059594': 'In stock at JLCPCB (1,850): no part lead time. Beyond stock: global sourcing 9-20 business days.',
 'C59873': 'In stock at JLCPCB (1,993): no part lead time. Beyond stock: global sourcing 9-20 business days.',
}
DEFAULT_LEAD = 'In stock at JLCPCB: no part lead time; only the PCBA lead time applies.'
rows = []
for b in bom:
    c = b['JLCPCB Part #']; j = jlc[c]; refs = b['Designator'].split(','); qty = len(refs)
    price = j['componentPrices'][0]['productPrice'] if j.get('componentPrices') else None
    alt, note = ALT.get(c, ('', ''))
    rows.append(dict(lcsc=c, mpn=j['componentModelEn'], brand=j['componentBrandEn'], pkg=j['componentSpecificationEn'], refs=b['Designator'], qty=qty,
                     stock=j['stockCount'], lib='Basic' if j['componentLibraryType']=='base' else 'Extended', price=price, alt=alt, note=note, lead=LEAD.get(c, DEFAULT_LEAD)))
BOARDS = 10
with open('rev4/order/STOCK_CHECK.csv', 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f); w.writerow(['JLCPCB Part #','MPN','Manufacturer','Package','Designators','Qty per board',f'Needed for {BOARDS} boards','JLCPCB stock (2026-10-13)','Library','Unit price USD (qty 1)','Status','Alternative','Note','Lead time / sourcing'])
    for r in rows:
        need = r['qty']*BOARDS; st = 'SHORT' if r['stock'] < need else ('LOW' if r['stock'] < 5*need else 'OK')
        w.writerow([r['lcsc'], r['mpn'], r['brand'], r['pkg'], r['refs'], r['qty'], need, r['stock'], r['lib'], r['price'], st, r['alt'], r['note'], r['lead']])
# ---- workbook
wb = Workbook(); F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='16324F'); body = Font(name=F, size=10)
thin = Side(style='thin', color='C5CCD3'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
def sheet(ws, headers, data, widths):
    ws.append(headers)
    for c in ws[1]: c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = border
    for row in data:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = body; c.alignment = Alignment(vertical='top', wrap_text=True); c.border = border
    for i, wdt in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = 'A2'
ws = wb.active; ws.title = 'BOM_SMT'
sheet(ws, ['Comment', 'Designator', 'Footprint', 'JLCPCB Part #'], [[b['Comment'], b['Designator'], b['Footprint'], b['JLCPCB Part #']] for b in bom], [62, 34, 18, 14])
ws2 = wb.create_sheet('Hand_Solder')
sheet(ws2, list(hand[0].keys()), [list(h.values()) for h in hand], [14, 44, 24, 40, 8, 5, 60])
ws3 = wb.create_sheet('Stock_Check')
ws3['A1'] = 'Boards to order'; ws3['B1'] = BOARDS; ws3['A1'].font = Font(name=F, bold=True); ws3['B1'].font = Font(name=F, bold=True, color='0000FF'); ws3['B1'].fill = PatternFill('solid', fgColor='FFFF00')
ws3['C1'] = 'Edit B1. Needed = qty per board x B1; Status: SHORT if stock < needed, LOW if stock < 5 x needed. Stock is JLCPCB assembly stock from the JLCPCB parts API, 2026-10-13.'; ws3['C1'].font = Font(name=F, italic=True, size=9)
ws3.append([])
headers = ['JLCPCB Part #','MPN','Manufacturer','Package','Designators','Qty per board','Needed','JLCPCB stock','Library','Unit price USD','Status','Alternative','Note','Lead time / sourcing']
ws3.append(headers)
for c in ws3[3]: c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = border
for i, r in enumerate(rows, start=4):
    ws3.append([r['lcsc'], r['mpn'], r['brand'], r['pkg'], r['refs'], r['qty'], f'=F{i}*$B$1', r['stock'], r['lib'], r['price'], f'=IF(H{i}<G{i},"SHORT",IF(H{i}<5*G{i},"LOW","OK"))', r['alt'], r['note'], r['lead']])
for row in ws3.iter_rows(min_row=4):
    for c in row: c.font = body; c.alignment = Alignment(vertical='top', wrap_text=True); c.border = border
for i, wdt in enumerate([13, 24, 20, 16, 30, 7, 8, 12, 10, 10, 8, 14, 60, 60], 1): ws3.column_dimensions[get_column_letter(i)].width = wdt
ws3.freeze_panes = 'A4'
wb.save('rev4/order/Rev4_JLCPCB_Order.xlsx'); print('xlsx written; stock rows', len(rows))
